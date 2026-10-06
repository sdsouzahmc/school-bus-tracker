"""Validated import of the 'School Bus Trip Dataset' workbook (Students, Crew, Locations, Trips, StudentTrips, ScanEvents,
Exceptions, Notifications) into the application database. Spec §2: validated import; errors are reported, nothing is half-loaded."""
import math
from datetime import date, datetime, timedelta

import openpyxl

import db

REQUIRED = {
    "Students": ["Student_ID", "Student_Name", "Batch", "Student_Group", "Class", "Bus_ID", "Route_ID", "Home_Stop_ID", "Guardian_ID",
                 "Guardian_Name", "Guardian_Email", "Authorized_Recipient_ID", "Badge_Token", "Student_Status"],
    "Crew": ["Crew_ID", "Crew_Name", "Role", "Assigned_Bus_ID", "Demo_Email", "Status"],
    "Locations": ["Location_ID", "Location_Name", "Location_Type", "Latitude", "Longitude"],
    "Trips": ["Trip_ID", "Service_Date", "Trip_No", "Bus_ID", "Route_ID", "Driver_ID", "Supervisor_ID", "Caretaker_ID", "Actual_Start",
              "Start_Lat", "Start_Lon", "Actual_End", "End_Lat", "End_Lon", "Trip_Status", "Sweep_Time", "Sweep_Performer_ID", "Student_Capacity"],
    "StudentTrips": ["Line_ID", "Trip_ID", "Student_ID", "Batch", "Student_Group", "Movement_Purpose", "Home_Stop_ID", "Entry_Event_ID",
                     "Entry_Time", "Entry_Lat", "Entry_Lon", "Exit_Event_ID", "Exit_Time", "Exit_Lat", "Exit_Lon", "Receiver_ID", "Handover_Type",
                     "Student_Status", "Capture_Method", "Reason"],
    "ScanEvents": ["Event_ID", "Trip_ID", "Student_ID", "Badge_Token", "Event_Type", "Capture_Time", "Server_Receipt_Time", "Location_ID",
                   "Latitude", "Longitude", "Accuracy_Metres", "GPS_Status", "Capture_Method", "Accepted", "Validation_Result", "Operator_ID",
                   "Device_ID", "Device_Sequence", "Capture_Connectivity", "Sync_Status"],
    "Exceptions": ["Exception_ID", "Trip_ID", "Student_ID", "Exception_Type", "Severity", "Triggered_Time", "Acknowledged_Time", "Owner_ID",
                   "Case_Status", "Action_or_Resolution", "Resolved_Time"],
    "Notifications": ["Notification_ID", "Event_ID", "Trip_ID", "Student_ID", "Guardian_ID", "Channel", "Notice_Type", "Queued_Time",
                      "Last_Attempt_Time", "Delivery_Status", "Attempt_Count", "Delayed_Update_Label", "Error_Detail"],
}
LINE_STATUS = {"Completed transport": "COMPLETED", "Declared absent": "ABSENT_DECLARED", "No show": "NO_SHOW",
               "Cancelled transport": "CANCELLED", "Onboard": "ONBOARD", "Expected": "EXPECTED", "Returned to school": "RETURNED"}
HANDOVER = {"School staff receipt": "SCHOOL_RECEIPT", "Authorized guardian release": "HOME_RELEASE", "Returned to school": "RETURNED_TO_SCHOOL"}
INC_KIND = {"Duplicate scan": "DUPLICATE_SCAN", "Wrong bus attempt": "WRONG_BUS", "Missing boarding": "NO_SHOW_PICKUP",
            "Recipient delayed": "RECIPIENT_DELAYED", "GPS unavailable": "GPS_UNAVAILABLE", "Missing check-out": "MISSING_CHECKOUT"}
NOTICE = {"Boarded": "BOARDED", "School receipt": "ARRIVED", "Authorized release": "RELEASED"}
ROLE = {"Driver": "driver", "Supervisor": "supervisor", "Care-taker": "caretaker", "Transport manager": "manager"}


SHIFT = timedelta(0)


def _ts(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return (v + SHIFT).strftime("%Y-%m-%d %H:%M:%S")
    return str(v)


def _rows(wb, sheet):
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        if r is None or all(v is None for v in r):
            continue
        out.append(dict(zip(header, r)))
    return header, out


def read_and_validate(path_or_file):
    """Returns (data, errors, warnings). data is None when errors exist."""
    errors, warnings = [], []
    try:
        wb = openpyxl.load_workbook(path_or_file, data_only=True, read_only=True)
    except Exception as e:      # noqa: BLE001
        return None, [f"Cannot open workbook: {e}"], []
    data = {}
    for sheet, cols in REQUIRED.items():
        if sheet not in wb.sheetnames:
            errors.append(f"Missing sheet '{sheet}'.")
            continue
        header, rows = _rows(wb, sheet)
        missing = [c for c in cols if c not in header]
        if missing:
            errors.append(f"Sheet '{sheet}' is missing column(s): {', '.join(missing)}.")
        data[sheet] = rows
    if errors:
        return None, errors, warnings

    def ids(sheet, col):
        return {r[col] for r in data[sheet]}

    def dup(sheet, col):
        seen, d = set(), set()
        for r in data[sheet]:
            (d if r[col] in seen else seen).add(r[col])
        if d:
            errors.append(f"{sheet}: duplicate {col} {sorted(d)[:5]}")

    for sheet, col in (("Students", "Student_ID"), ("Crew", "Crew_ID"), ("Locations", "Location_ID"), ("Trips", "Trip_ID"),
                       ("StudentTrips", "Line_ID"), ("ScanEvents", "Event_ID"), ("Students", "Badge_Token")):
        dup(sheet, col)
    loc, crew, stu, trips, ev = (ids("Locations", "Location_ID"), ids("Crew", "Crew_ID"), ids("Students", "Student_ID"),
                                 ids("Trips", "Trip_ID"), ids("ScanEvents", "Event_ID"))
    if not any(r["Location_Type"] == "School" for r in data["Locations"]):
        errors.append("Locations: no row with Location_Type = School.")
    for r in data["Students"]:
        if r["Home_Stop_ID"] not in loc:
            errors.append(f"Students {r['Student_ID']}: unknown Home_Stop_ID {r['Home_Stop_ID']}")
        if r["Batch"] not in (1, 2):
            errors.append(f"Students {r['Student_ID']}: Batch must be 1 or 2")
        if r["Student_Group"] not in ("Junior", "Senior"):
            errors.append(f"Students {r['Student_ID']}: Student_Group must be Junior or Senior")
        if not r["Badge_Token"]:
            warnings.append(f"Students {r['Student_ID']}: no badge token — only manual entry possible")
    for r in data["Trips"]:
        for c in ("Driver_ID", "Supervisor_ID", "Caretaker_ID"):
            if r[c] not in crew:
                errors.append(f"Trips {r['Trip_ID']}: unknown {c} {r[c]}")
        if r["Trip_No"] not in db.TRIPS:
            errors.append(f"Trips {r['Trip_ID']}: Trip_No must be 1–5")
    for r in data["StudentTrips"]:
        if r["Trip_ID"] not in trips:
            errors.append(f"StudentTrips {r['Line_ID']}: unknown Trip_ID {r['Trip_ID']}")
        if r["Student_ID"] not in stu:
            errors.append(f"StudentTrips {r['Line_ID']}: unknown Student_ID {r['Student_ID']}")
        if r["Student_Status"] not in LINE_STATUS:
            errors.append(f"StudentTrips {r['Line_ID']}: unknown Student_Status '{r['Student_Status']}'")
        for c in ("Entry_Event_ID", "Exit_Event_ID"):
            if r[c] and r[c] not in ev:
                errors.append(f"StudentTrips {r['Line_ID']}: {c} {r[c]} not found in ScanEvents")
        if r["Exit_Event_ID"] and not r["Entry_Event_ID"]:
            errors.append(f"StudentTrips {r['Line_ID']}: exit without entry")
    for r in data["ScanEvents"]:
        if r["Trip_ID"] not in trips:
            errors.append(f"ScanEvents {r['Event_ID']}: unknown Trip_ID")
        if r["Location_ID"] and r["Location_ID"] not in loc:
            errors.append(f"ScanEvents {r['Event_ID']}: unknown Location_ID {r['Location_ID']}")
    if len(errors) > 50:
        errors = errors[:50] + [f"… and {len(errors) - 50} more"]
    return (None if errors else data), errors, warnings


def load(data, source_name="School_Bus_Trip_Dataset.xlsx", shift_to_past=True):
    """Replace the database with the workbook's data. If the sample lies in the future (relative to today), move it back by whole
    weeks so weekdays are kept and 'today' in the demo comes after the history."""
    global SHIFT
    days = [r["Service_Date"] for r in data["Trips"] if isinstance(r["Service_Date"], datetime)]
    SHIFT = timedelta(0)
    note = ""
    if shift_to_past and days:
        first, last = min(days).date(), max(days).date()
        today = date.fromisoformat(db.today())
        if last >= today:
            weeks = math.ceil(((last - today).days + 1) / 7)
            SHIFT = timedelta(weeks=-weeks)
            note = (f"; dates moved back {weeks} week(s): {first:%d %b}–{last:%d %b} → {(first + SHIFT):%d %b}–{(last + SHIFT):%d %b %Y}")
    db.init_db(reset=True, seed_data=False)
    con = db.connect()
    try:
        c = con.cursor()
        for x in data["Crew"]:
            x["Crew_Name"] = str(x["Crew_Name"]).replace(" Demo", "")
        school = next(r for r in data["Locations"] if r["Location_Type"] == "School")
        for k, v in (("school_lat", school["Latitude"]), ("school_lon", school["Longitude"]), ("school_name", db.SCHOOL_NAME if "demo" in str(school["Location_Name"]).lower() else school["Location_Name"]),
                     ("data_source", f"{source_name} ({len(data['Students'])} students, "
                                     f"{len({r['Bus_ID'] for r in data['Trips']})} buses{note})")):
            c.execute("INSERT OR REPLACE INTO settings VALUES (?,?)", (k, str(v)))
        db._POS.clear()

        # staff & users
        staff = {}
        for r in data["Crew"]:
            staff[r["Crew_ID"]] = c.execute("INSERT INTO staff(name, role, phone) VALUES (?,?,?)",
                                            (r["Crew_Name"], ROLE.get(r["Role"], r["Role"].lower()), r["Demo_Email"])).lastrowid
        recv_ids = sorted({r["Receiver_ID"] for r in data["StudentTrips"] if r["Receiver_ID"] and not str(r["Receiver_ID"]).startswith("G")})
        for rid in recv_ids or ["SCHOOL_RECEIVER_01"]:
            staff[rid] = c.execute("INSERT INTO staff(name, role, phone) VALUES (?,?,?)",
                                   (rid.replace("_", " ").title().replace("School Receiver", "School Receiver"), "receiving", "")).lastrowid

        def user(u, pin, role, display, staff_id=None, guardian_id=None):
            c.execute("INSERT OR REPLACE INTO users(username, pin_hash, role, staff_id, guardian_id, display_name) VALUES (?,?,?,?,?,?)",
                      (u, db.pin_hash(pin), role, staff_id, guardian_id, display))

        # buses (route -> bus), stops, devices
        bus_ids, route_bus = {}, {}
        for r in sorted(data["Trips"], key=lambda r: r["Bus_ID"]):
            if r["Bus_ID"] in bus_ids:
                continue
            n = len(bus_ids) + 1
            bid = c.execute("INSERT INTO buses(bus_no, plate, capacity, fvts_vehicle, route_name, driver_id, supervisor_id, caretaker_id) "
                            "VALUES (?,?,?,?,?,?,?,?)", (r["Bus_ID"], "", int(r["Student_Capacity"] or 40), r["Bus_ID"], f"Route {r['Route_ID']}",
                                                         staff[r["Driver_ID"]], staff[r["Supervisor_ID"]], staff[r["Caretaker_ID"]])).lastrowid
            bus_ids[r["Bus_ID"]] = bid
            route_bus[r["Route_ID"]] = bid
            for role, col, pre in (("driver", "Driver_ID", "drv"), ("supervisor", "Supervisor_ID", "sup"), ("caretaker", "Caretaker_ID", "care")):
                c.execute("UPDATE staff SET bus_id=? WHERE id=?", (bid, staff[r[col]]))
                cr = next(x for x in data["Crew"] if x["Crew_ID"] == r[col])
                user(f"{pre}{n:02d}", "1234", role, cr["Crew_Name"], staff[r[col]])
        dev_ids = {}
        for r in data["ScanEvents"]:
            if r["Device_ID"] and r["Device_ID"] not in dev_ids:
                trip = next(t for t in data["Trips"] if t["Trip_ID"] == r["Trip_ID"])
                dev_ids[r["Device_ID"]] = c.execute("INSERT INTO devices(name, bus_id, status, registered_at) VALUES (?,?,?,?)",
                                                    (f"{r['Device_ID']} ({trip['Bus_ID']} tablet)", bus_ids[trip["Bus_ID"]], "ACTIVE",
                                                     db.ts())).lastrowid
        for b, bid in bus_ids.items():
            if not c.execute("SELECT 1 FROM devices WHERE bus_id=?", (bid,)).fetchone():
                c.execute("INSERT INTO devices(name, bus_id, status, registered_at) VALUES (?,?,?,?)", (f"{b} tablet", bid, "ACTIVE", db.ts()))
        stop_ids = {school["Location_ID"]: db.SCHOOL_STOP}
        for r in data["Locations"]:
            if r["Location_Type"] == "Home stop":
                route = r["Location_ID"].split("S")[0]
                seq = int(r["Location_ID"].split("S")[-1]) if r["Location_ID"].split("S")[-1].isdigit() else 0
                stop_ids[r["Location_ID"]] = c.execute("INSERT INTO stops(bus_id, seq, name, area, lat, lon) VALUES (?,?,?,?,?,?)",
                                                       (route_bus.get(route), seq, r["Location_Name"], route, r["Latitude"], r["Longitude"])).lastrowid
        loc_name = {r["Location_ID"]: r["Location_Name"] for r in data["Locations"]}

        # guardians, students, recipients, badges
        g_ids, stu_ids = {}, {}
        for r in data["Students"]:
            if r["Guardian_ID"] not in g_ids:
                g_ids[r["Guardian_ID"]] = c.execute("INSERT INTO guardians(name, phone, email, relation, verified) VALUES (?,?,?,?,1)",
                                                    (r["Guardian_Name"], "", r["Guardian_Email"], "Guardian")).lastrowid
            sid = c.execute("INSERT INTO students(code, name, class_name, section, batch, grp, bus_id, stop_id, active) VALUES (?,?,?,?,?,?,?,?,?)",
                            (r["Student_ID"], r["Student_Name"], r["Class"], "", int(r["Batch"]), r["Student_Group"],
                             bus_ids.get(r["Bus_ID"]), stop_ids.get(r["Home_Stop_ID"]), 1 if r["Student_Status"] == "Active" else 0)).lastrowid
            stu_ids[r["Student_ID"]] = sid
            c.execute("INSERT INTO student_guardians VALUES (?,?,1)", (sid, g_ids[r["Guardian_ID"]]))
            if r["Authorized_Recipient_ID"]:
                rec = next((s for s in data["Students"] if s["Guardian_ID"] == r["Authorized_Recipient_ID"]), r)
                c.execute("INSERT INTO recipients(student_id, name, relation, phone, id_number) VALUES (?,?,?,?,?)",
                          (sid, rec["Guardian_Name"], "Guardian", "", r["Authorized_Recipient_ID"]))
            if r["Badge_Token"]:
                c.execute("INSERT INTO badges(code, student_id, status, issued_at) VALUES (?,?,?,?)",
                          (str(r["Badge_Token"]).upper(), sid, "ACTIVE", db.ts()))
        for i, (gcode, gid) in enumerate(sorted(g_ids.items()), start=1):
            name = c.execute("SELECT name FROM guardians WHERE id=?", (gid,)).fetchone()[0]
            user(f"g{i:03d}", "1234", "guardian", name, None, gid)
        g_name = {r["Guardian_ID"]: r["Guardian_Name"] for r in data["Students"]}

        # trips
        run_ids = {}
        for r in data["Trips"]:
            status = {"Completed": "CLOSED", "Reconciling": "IN_PROGRESS", "Pending sync": "PENDING_SYNC", "Cancelled": "CLOSED"}.get(r["Trip_Status"], "IN_PROGRESS")
            closed = status == "CLOSED"
            sweep = None
            if r["Sweep_Performer_ID"] and r["Sweep_Time"]:
                sweep = next(x["Crew_Name"] for x in data["Crew"] if x["Crew_ID"] == r["Sweep_Performer_ID"])
            run_ids[r["Trip_ID"]] = c.execute(
                """INSERT INTO trip_runs(bus_id, date, trip_no, status, start_ts, start_lat, start_lon, end_ts, end_lat, end_lon, driver_id,
                   supervisor_id, caretaker_id, final_driver_id, final_supervisor_id, final_caretaker_id, sweep_by, sweep_ts, started_by, closed_by,
                   device_id, cancelled) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (bus_ids[r["Bus_ID"]], _ts(r["Service_Date"])[:10], int(r["Trip_No"]), status, _ts(r["Actual_Start"]), r["Start_Lat"], r["Start_Lon"],
                 _ts(r["Actual_End"]), r["End_Lat"] if r["Actual_End"] else None, r["End_Lon"] if r["Actual_End"] else None,
                 staff[r["Driver_ID"]], staff[r["Supervisor_ID"]], staff[r["Caretaker_ID"]],
                 staff[r["Driver_ID"]] if closed else None, staff[r["Supervisor_ID"]] if closed else None, staff[r["Caretaker_ID"]] if closed else None,
                 sweep, _ts(r["Sweep_Time"]) if sweep else None, r["Supervisor_ID"], r["Supervisor_ID"] if closed else None,
                 None, 1 if r["Trip_Status"] == "Cancelled" else 0)).lastrowid

        # student trip lines
        line_of_event, man_ids = {}, {}
        line_reason = {}
        for r in data["StudentTrips"]:
            rid = r["Receiver_ID"]
            ho = HANDOVER.get(r["Handover_Type"])
            mid = c.execute(
                """INSERT INTO manifest(run_id, student_id, purpose, batch, grp, stop_id, status, entry_ts, entry_lat, entry_lon, exit_ts, exit_lat,
                   exit_lon, handover, recipient_name, recipient_relation, received_by, note) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (run_ids[r["Trip_ID"]], stu_ids[r["Student_ID"]], "PICKUP" if r["Movement_Purpose"] == "Pickup" else "DROPOFF", int(r["Batch"]),
                 r["Student_Group"], stop_ids.get(r["Home_Stop_ID"]), LINE_STATUS[r["Student_Status"]], _ts(r["Entry_Time"]), r["Entry_Lat"],
                 r["Entry_Lon"], _ts(r["Exit_Time"]), r["Exit_Lat"], r["Exit_Lon"], ho,
                 g_name.get(rid) if ho == "HOME_RELEASE" else None, "Guardian" if ho == "HOME_RELEASE" else None,
                 rid.replace("_", " ").title() if ho == "SCHOOL_RECEIPT" and rid else None, r["Reason"])).lastrowid
            man_ids[r["Line_ID"]] = mid
            for e in (r["Entry_Event_ID"], r["Exit_Event_ID"]):
                if e:
                    line_of_event[e] = mid
                    line_reason[e] = r["Reason"]
            if r["Student_Status"] == "Declared absent":
                d = c.execute("SELECT date, trip_no FROM trip_runs WHERE id=?", (run_ids[r["Trip_ID"]],)).fetchone()
                if not c.execute("SELECT 1 FROM absences WHERE student_id=? AND date=?", (stu_ids[r["Student_ID"]], d[0])).fetchone():
                    c.execute("INSERT INTO absences(student_id, date, trip_no, reason, declared_by, source, status, created_at) VALUES (?,?,?,?,?,?,?,?)",
                              (stu_ids[r["Student_ID"]], d[0], None, r["Reason"] or "Declared absent", "guardian", "GUARDIAN", "APPROVED", d[0] + " 05:00:00"))

        # scan events
        crew_name = {x["Crew_ID"]: x["Crew_Name"] for x in data["Crew"]}
        trip_row = {t["Trip_ID"]: t for t in data["Trips"]}
        for r in data["ScanEvents"]:
            accepted = int(r["Accepted"] or 0) == 1
            manual = r["Capture_Method"] == "Manual"
            t = trip_row[r["Trip_ID"]]
            c.execute(
                """INSERT INTO events(client_uuid, run_id, manifest_id, student_id, badge_code, kind, result, reason, captured_ts, received_ts, lat, lon,
                   accuracy, gps_ok, stop_id, operator, device_id, method, manual_reason, verified_by, sync_state, clock_flag, connectivity, device_seq)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (r["Event_ID"], run_ids[r["Trip_ID"]], line_of_event.get(r["Event_ID"]), stu_ids.get(r["Student_ID"]),
                 str(r["Badge_Token"]).upper() if r["Badge_Token"] else None, r["Event_Type"], "ACCEPTED" if accepted else "REJECTED",
                 None if r["Validation_Result"] == "Accepted" else r["Validation_Result"], _ts(r["Capture_Time"]), _ts(r["Server_Receipt_Time"]),
                 r["Latitude"], r["Longitude"], r["Accuracy_Metres"], 0 if r["GPS_Status"] != "Valid" else 1, stop_ids.get(r["Location_ID"]),
                 crew_name.get(r["Operator_ID"], r["Operator_ID"]), dev_ids.get(r["Device_ID"]), "MANUAL" if manual else "QR",
                 line_reason.get(r["Event_ID"]) if manual else None, crew_name.get(t["Caretaker_ID"]) if manual else None,
                 "SYNCED" if r["Sync_Status"] == "Synced" else "PENDING", None, r["Capture_Connectivity"] or "Online", r["Device_Sequence"]))

        # exceptions -> incidents
        mgr = {x["Crew_ID"]: x["Crew_Name"] for x in data["Crew"]}
        for r in data["Exceptions"]:
            run = run_ids.get(r["Trip_ID"])
            bus = c.execute("SELECT bus_id FROM trip_runs WHERE id=?", (run,)).fetchone()
            sname = next((s["Student_Name"] for s in data["Students"] if s["Student_ID"] == r["Student_ID"]), "")
            c.execute("""INSERT INTO incidents(ts, run_id, bus_id, student_id, kind, severity, details, status, assigned_to, resolution, resolved_ts, ack_ts)
                         VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                      (_ts(r["Triggered_Time"]), run, bus[0] if bus else None, stu_ids.get(r["Student_ID"]),
                       INC_KIND.get(r["Exception_Type"], r["Exception_Type"].upper().replace(" ", "_")), str(r["Severity"]).upper(),
                       f"{r['Exception_Type']}: {sname} ({r['Student_ID']}) — {trip_row[r['Trip_ID']]['Bus_ID']} Trip {trip_row[r['Trip_ID']]['Trip_No']} "
                       f"on {_ts(trip_row[r['Trip_ID']]['Service_Date'])[:10]} (source {r['Exception_ID']})",
                       "RESOLVED" if r["Case_Status"] == "Resolved" else ("ASSIGNED" if r["Acknowledged_Time"] else "OPEN"),
                       mgr.get(r["Owner_ID"], r["Owner_ID"]), r["Action_or_Resolution"], _ts(r["Resolved_Time"]), _ts(r["Acknowledged_Time"])))

        # notifications
        ev_row = {e["Event_ID"]: e for e in data["ScanEvents"]}
        stu_name = {s["Student_ID"]: s["Student_Name"] for s in data["Students"]}
        for r in data["Notifications"]:
            t = trip_row.get(r["Trip_ID"], {})
            e = ev_row.get(r["Event_ID"], {})
            hm = (_ts(e.get("Capture_Time")) or "")[11:16]
            where = loc_name.get(e.get("Location_ID"), "")
            msg = {"Boarded": f"{stu_name.get(r['Student_ID'])} boarded {t.get('Bus_ID')} Trip {t.get('Trip_No')} at {where} at {hm}.",
                   "School receipt": f"{stu_name.get(r['Student_ID'])} arrived at school on {t.get('Bus_ID')} Trip {t.get('Trip_No')} at {hm}.",
                   "Authorized release": f"{stu_name.get(r['Student_ID'])} was dropped at {where} at {hm} and handed to an authorized guardian."
                   }.get(r["Notice_Type"], f"{r['Notice_Type']}: {stu_name.get(r['Student_ID'])} at {hm}")
            sent = r["Delivery_Status"] == "Sent"
            c.execute("""INSERT INTO notifications(ts, guardian_id, student_id, run_id, channel, kind, message, status, attempts, last_error,
                         delayed_offline, sent_ts) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                      (_ts(r["Queued_Time"]), g_ids.get(r["Guardian_ID"]), stu_ids.get(r["Student_ID"]), run_ids.get(r["Trip_ID"]),
                       "EMAIL" if str(r["Channel"]).lower() == "email" else "IN_APP", NOTICE.get(r["Notice_Type"], r["Notice_Type"].upper()), msg,
                       "SENT" if sent else "FAILED", int(r["Attempt_Count"] or 1), r["Error_Detail"],
                       1 if r["Delayed_Update_Label"] else 0, _ts(r["Last_Attempt_Time"]) if sent else None))

        # office users
        mid = next((staff[x["Crew_ID"]] for x in data["Crew"] if ROLE.get(x["Role"]) == "manager"), None)
        aid = c.execute("INSERT INTO staff(name, role) VALUES ('System Administrator','admin')").lastrowid
        did = c.execute("INSERT INTO staff(name, role) VALUES ('Dispatcher','dispatcher')").lastrowid
        user("admin", "admin123", "admin", "System Administrator", aid)
        user("manager", "manager123", "manager", next((x["Crew_Name"] for x in data["Crew"] if ROLE.get(x["Role"]) == "manager"), "Manager"), mid)
        user("dispatcher", "1234", "dispatcher", "Dispatcher", did)
        for i, rid in enumerate(recv_ids or ["SCHOOL_RECEIVER_01"], start=1):
            user(f"recv{i}", "1234", "receiving", rid.replace("_", " ").title(), staff[rid])
        c.execute("INSERT INTO audit(ts, username, role, action, entity, entity_id, details) VALUES (?,?,?,?,?,?,?)",
                  (db.ts(), "system", "admin", "DATA_IMPORT", "workbook", source_name,
                   f"{len(data['Students'])} students, {len(data['Trips'])} trips, {len(data['ScanEvents'])} scan events"))
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()
    return {"students": len(data["Students"]), "buses": len(bus_ids), "trips": len(data["Trips"]), "lines": len(data["StudentTrips"]),
            "events": len(data["ScanEvents"]), "incidents": len(data["Exceptions"]), "notifications": len(data["Notifications"])}


def import_workbook(path_or_file, source_name="School_Bus_Trip_Dataset.xlsx", shift_to_past=True):
    data, errors, warnings = read_and_validate(path_or_file)
    if errors:
        return None, errors, warnings
    return load(data, source_name, shift_to_past), [], warnings
